from docx import Document
import os

# Track current working file
current_file = None
document = None

def create_word_file(filename):
    global current_file, document
    if not filename.endswith(".docx"):
        filename += ".docx"

    document = Document()
    current_file = os.path.join(os.getcwd(), filename)
    document.save(current_file)
    return f"Created a Word file named {filename}."

def write_to_file(text):
    global current_file, document
    if not current_file or not document:
        return "No file is open to write to."

    document.add_paragraph(text)
    document.save(current_file)
    return "Written to the file."

def add_bullet_point(text):
    global current_file, document
    if not current_file or not document:
        return "No file is open to write to."

    document.add_paragraph(text, style='List Bullet')
    document.save(current_file)
    return "Bullet point added."

def close_file():
    global current_file, document
    if not current_file:
        return "No file is currently open."
    
    document.save(current_file)
    name = os.path.basename(current_file)
    current_file = None
    document = None
    return f"File {name} saved and closed."
